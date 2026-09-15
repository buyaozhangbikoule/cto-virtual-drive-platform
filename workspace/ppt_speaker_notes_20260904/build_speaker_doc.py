from __future__ import annotations

import json
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"C:\Users\24739\Documents\Codex\2026-08-10\new-chat")
WORK = ROOT / "ppt_speaker_notes_20260904"
OUTPUT = ROOT / "CTO超声换能器第一月组会_逐页讲稿与答辩备忘_2026-09-04.docx"
CONTENT = json.loads((WORK / "speaker_content.json").read_text(encoding="utf-8"))

NAVY = "17365D"
BLUE = "2E74B5"
LIGHT_BLUE = "E8EEF5"
PALE_YELLOW = "FFF4DD"
PALE_RED = "FDECEC"
GRAY = "666666"
LIGHT_GRAY = "F3F5F7"
WHITE = "FFFFFF"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=140, bottom=100, end=140) -> None:
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, size=11, bold=False, color="000000", italic=False) -> None:
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "微软雅黑")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor.from_string(color)


def set_para_format(p, before=0, after=6, line=1.25, keep=False) -> None:
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = line
    p.paragraph_format.keep_with_next = keep


def add_labeled_para(doc, label: str, text: str, color=BLUE, fill=None, keep=False):
    if fill:
        table = doc.add_table(rows=1, cols=1)
        table.autofit = False
        table.columns[0].width = Inches(6.5)
        cell = table.cell(0, 0)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(cell, fill)
        set_cell_margins(cell)
        p = cell.paragraphs[0]
    else:
        p = doc.add_paragraph()
    set_para_format(p, after=5, line=1.2, keep=keep)
    r = p.add_run(label)
    set_font(r, 10.5, True, color)
    r = p.add_run(text)
    set_font(r, 10.5, False, "1F1F1F")
    return p


def add_field(paragraph, field_code: str) -> None:
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = field_code
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr, fld_char2])


doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.72)
section.bottom_margin = Inches(0.68)
section.left_margin = Inches(0.88)
section.right_margin = Inches(0.88)
section.header_distance = Inches(0.32)
section.footer_distance = Inches(0.3)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Calibri"
normal._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
normal.font.size = Pt(11)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.25

for name, size, color, before, after in [
    ("Title", 23, NAVY, 0, 8),
    ("Heading 1", 16, BLUE, 18, 10),
    ("Heading 2", 13, BLUE, 14, 7),
    ("Heading 3", 12, "1F4D78", 10, 5),
]:
    style = styles[name]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "微软雅黑")
    style.font.size = Pt(size)
    style.font.bold = True
    style.font.color.rgb = RGBColor.from_string(color)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

header = section.header.paragraphs[0]
header.alignment = WD_ALIGN_PARAGRAPH.LEFT
set_para_format(header, after=0, line=1.0)
r = header.add_run("CTO换能器虚拟电学平台  |  第一月组会逐页讲稿")
set_font(r, 8.5, False, GRAY)

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
set_para_format(footer, after=0, line=1.0)
r = footer.add_run("内部汇报材料  ·  第 ")
set_font(r, 8.5, False, GRAY)
add_field(footer, "PAGE")
r = footer.add_run(" 页 / 共 ")
set_font(r, 8.5, False, GRAY)
add_field(footer, "NUMPAGES")
r = footer.add_run(" 页")
set_font(r, 8.5, False, GRAY)

# Memo-masthead cover.
p = doc.add_paragraph()
set_para_format(p, before=18, after=4)
r = p.add_run("研究生组会汇报讲稿")
set_font(r, 12, True, BLUE)
p = doc.add_paragraph(style="Title")
r = p.add_run("CTO超声换能器虚拟阻抗建模\n与自适应驱动控制")
set_font(r, 23, True, NAVY)
p = doc.add_paragraph()
set_para_format(p, after=18)
r = p.add_run("第一月 · 28页逐页讲解、口径审计与答辩备忘")
set_font(r, 14, False, "373737")

meta = [
    ("配套PPT", "CTO超声换能器第一月教授组会汇报_逐页讲稿版_2026-09-04.pptx"),
    ("建议时长", "正文第1–18页约14–15分钟；第19–28页仅在答疑时调用"),
    ("证据边界", "假设小信号BVD与v2.1匹配网络；无实测、无结构声场、无临床结论"),
    ("不确定性标记", "所有弱因果或尚未验证的延伸均显式标注【模糊推测】"),
    ("版本日期", "2026-09-04"),
]
for label, value in meta:
    p = doc.add_paragraph()
    set_para_format(p, after=3, line=1.15)
    r = p.add_run(f"{label}：")
    set_font(r, 10.5, True, NAVY)
    r = p.add_run(value)
    set_font(r, 10.5, False, "1F1F1F")

p = doc.add_paragraph()
set_para_format(p, before=14, after=0)
p_pr = p._p.get_or_add_pPr()
p_bdr = OxmlElement("w:pBdr")
bottom = OxmlElement("w:bottom")
bottom.set(qn("w:val"), "single")
bottom.set(qn("w:sz"), "18")
bottom.set(qn("w:color"), NAVY)
p_bdr.append(bottom)
p_pr.append(p_bdr)

add_labeled_para(doc, "核心结论：", "本月已完成模型内电学与带约束控制闭环；外部有效性仍需真实阻抗、功率和声场实验。", fill=LIGHT_BLUE)
add_labeled_para(doc, "汇报原则：", "先报证据，再给解释；先说测点和单位，再说趋势；不把实现一致性误写成器件真实性。", fill=LIGHT_GRAY)

doc.add_page_break()
doc.add_heading("汇报路线与时间控制", level=1)
add_labeled_para(doc, "开场主线：", "物理冲突审计 → 双线验证 → 匹配 → 单因素归因 → 三准则 → 保护与控制 → 边界和下一步。", fill=LIGHT_BLUE)
p = doc.add_paragraph()
set_para_format(p, after=8)
r = p.add_run("正文页（建议实际讲）：")
set_font(r, 11, True, NAVY)
for item in CONTENT[:18]:
    p = doc.add_paragraph(style="List Number")
    set_para_format(p, after=2, line=1.1)
    r = p.add_run(f"第{item['slide']}页｜{item['title']}（{item['duration']}）")
    set_font(r, 9.5, False, "222222")
p = doc.add_paragraph()
set_para_format(p, before=8, after=4)
r = p.add_run("备份页：")
set_font(r, 11, True, NAVY)
r = p.add_run("第19–28页只用于回答参数来源、扰动复现、相位定义、证据链和COMSOL缺口。")
set_font(r, 10.5, False, "222222")
add_labeled_para(doc, "教授组会提醒：", "表格里的0%不是“恢复到基准绝对功率”；LTspice/Python误差不是模型对实物的误差；保护阈值不是临床安全阈值。", fill=PALE_YELLOW)

for item in CONTENT:
    doc.add_page_break()
    section_label = "正文" if item["slide"] <= 18 else "备份 / 答疑"
    p = doc.add_paragraph()
    set_para_format(p, before=0, after=3)
    r = p.add_run(f"第 {item['slide']:02d} 页  ·  {section_label}  ·  建议 {item['duration']}")
    set_font(r, 9.5, True, BLUE)
    doc.add_heading(item["title"], level=1)
    add_labeled_para(doc, "本页结论：", item["conclusion"], fill=LIGHT_BLUE)
    doc.add_heading("逐字讲稿", level=2)
    p = doc.add_paragraph()
    set_para_format(p, after=7, line=1.25)
    r = p.add_run(item["script"])
    set_font(r, 11, False, "1F1F1F")
    add_labeled_para(doc, "看图顺序：", item["read_order"], color=BLUE)
    add_labeled_para(doc, "口径与限制：", item["caveat"], color="A94442", fill=PALE_RED)
    doc.add_heading("可能追问与回答", level=2)
    add_labeled_para(doc, "问：", item["question"], color=NAVY, keep=True)
    add_labeled_para(doc, "答：", item["answer"], color=BLUE)
    fuzzy_fill = PALE_YELLOW if item["fuzzy"] != "无。" else LIGHT_GRAY
    add_labeled_para(doc, "模糊推测：", item["fuzzy"], color="B36B00", fill=fuzzy_fill)
    p = doc.add_paragraph()
    set_para_format(p, before=4, after=0, line=1.05)
    r = p.add_run("证据路径：")
    set_font(r, 8.5, True, GRAY)
    for source in item["sources"]:
        full = source if ":/" in source else str(ROOT / source)
        r = p.add_run(f"\n• {full}")
        set_font(r, 8.2, False, GRAY)

doc.add_page_break()
doc.add_heading("答辩时必须守住的五条边界", level=1)
for text in [
    "模型参数是“假设基线值”，不是实测换能器参数。",
    "Python/LTspice一致性证明实现一致，不证明模型外部有效。",
    "P_Rm是运动支路电学有功功率代理，不等于声功率、声压或临床剂量。",
    "恢复时间是20 ms离散准稳态样本定义，不等于硬件瞬态时间。",
    "COMSOL/ANSYS尚未运行；接口PASS不等于声场PASS。",
]:
    p = doc.add_paragraph(style="List Bullet")
    set_para_format(p, after=6)
    r = p.add_run(text)
    set_font(r, 11, False, "1F1F1F")
add_labeled_para(doc, "建议结束语：", "本月最重要的产出不是某一个最优频率，而是一条能够暴露参数冲突、区分测点口径、复现实验场景并审计控制约束的电学仿真链。下一阶段以真实复阻抗拟合和结构声场验证为优先。", fill=LIGHT_BLUE)

doc.core_properties.title = "CTO超声换能器第一月组会逐页讲稿与答辩备忘"
doc.core_properties.subject = "假设BVD模型、电学匹配、扰动扫描、带约束控制与证据边界"
doc.core_properties.author = "Codex协助整理；数据来源为本地项目仿真输出"
doc.core_properties.keywords = "BVD, LTspice, Python, L-match, P&O, phase control, group meeting"
doc.save(OUTPUT)
print(OUTPUT)
