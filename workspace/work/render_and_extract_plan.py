from pathlib import Path
from zipfile import ZipFile

import pypdfium2 as pdfium
from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P

DOCX = Path(r"D:\微信文件\WeChat Files\wxid_f86c1kdajuft12\FileStorage\File\2026-08\CTO超声换能器虚拟阻抗建模与自适应驱动控制_第一月计划书.docx")
PDF = Path(r"C:\Users\24739\Documents\Codex\2026-08-10\new-chat\work\plan_docx_render\计划书_Word导出.pdf")
OUT = Path(r"C:\Users\24739\Documents\Codex\2026-08-10\new-chat\work\plan_docx_render")
OUT.mkdir(parents=True, exist_ok=True)


pdf = pdfium.PdfDocument(str(PDF))
for i in range(len(pdf)):
    page = pdf[i]
    bitmap = page.render(scale=2.0)
    bitmap.to_pil().save(OUT / f"page-{i + 1}.png")


doc = Document(str(DOCX))
body = doc.element.body
lines = []
paragraph_index = 0
table_index = 0
for child in body.iterchildren():
    if isinstance(child, CT_P):
        paragraph_index += 1
        p = Paragraph(child, doc)
        text = p.text.strip()
        if text:
            lines.append(f"[P{paragraph_index} style={p.style.name}] {text}")
    elif isinstance(child, CT_Tbl):
        table_index += 1
        table = Table(child, doc)
        lines.append(f"\n[TABLE {table_index} rows={len(table.rows)} cols={len(table.columns)}]")
        for r, row in enumerate(table.rows, 1):
            cells = [" ".join(cell.text.split()) for cell in row.cells]
            lines.append(f"R{r}: " + " | ".join(cells))
        lines.append(f"[END TABLE {table_index}]\n")

for section_no, section in enumerate(doc.sections, 1):
    header = " ".join(p.text.strip() for p in section.header.paragraphs if p.text.strip())
    footer = " ".join(p.text.strip() for p in section.footer.paragraphs if p.text.strip())
    if header:
        lines.append(f"[SECTION {section_no} HEADER] {header}")
    if footer:
        lines.append(f"[SECTION {section_no} FOOTER] {footer}")

with ZipFile(DOCX) as zf:
    names = set(zf.namelist())
    has_comments = "word/comments.xml" in names
    document_xml = zf.read("word/document.xml")
    has_insertions = b"<w:ins" in document_xml
    has_deletions = b"<w:del" in document_xml

lines.insert(0, f"pages={len(pdf)} tables={len(doc.tables)} paragraphs={len(doc.paragraphs)} comments={has_comments} insertions={has_insertions} deletions={has_deletions}")
(OUT / "计划书完整结构提取.txt").write_text("\n".join(lines), encoding="utf-8")
print(lines[0])
